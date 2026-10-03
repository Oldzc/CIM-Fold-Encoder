# -*- coding: utf-8 -*-
"""检查生成的报告：结构、残留标记、旧报告引用"""
import json
import re
import sys
from docx import Document

DOC = sys.argv[1]
CHECKS = sys.argv[2]

d = json.load(open(CHECKS, encoding='utf-8'))
print('结构检查:', d['verdict'])
print('段落数:', d['summary']['paragraphs'], ' 表格数:', len(d['summary']['tables']))
for c in d['checks']:
    if c['id'] == 'contains':
        print('  contains:', c['status'], c.get('text'))

doc = Document(DOC)
bad = [p.text[:70] for p in doc.paragraphs if '**' in p.text]
for t in doc.tables:
    for r in t.rows:
        for c in r.cells:
            if '**' in c.text:
                bad.append(c.text[:60])
print('残留 ** 标记:', len(bad))
for b in bad[:5]:
    print('   ', b)

print('旧报告引用检查:')
for k in ['首版', '原报告', '原工程', '补充报告', '偏离', '更正', '重写版']:
    n = sum(1 for p in doc.paragraphs if k in p.text)
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                if k in c.text:
                    n += 1
    flag = 'OK' if n == 0 else '<<< 仍存在'
    print('  %-8s %d 次  %s' % (k, n, flag))

print('章节:')
for p in doc.paragraphs:
    if re.match(r'^[一二三四五六七八九十]、', p.text.strip()):
        print('  ', p.text.strip())
